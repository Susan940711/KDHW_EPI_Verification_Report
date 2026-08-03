from io import BytesIO

import openpyxl
import pandas as pd


def load_excel_preview(uploaded_file) -> dict:
    workbook = openpyxl.load_workbook(BytesIO(uploaded_file.getvalue()), data_only=True)
    sheet_names = workbook.sheetnames
    first_sheet = workbook[sheet_names[0]]
    rows = list(first_sheet.iter_rows(values_only=True))

    if not rows:
        preview = pd.DataFrame()
    else:
        preview = pd.DataFrame(rows[1:], columns=rows[0])

    return {"sheet_names": sheet_names, "preview": preview}
